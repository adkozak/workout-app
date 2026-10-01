"""Snapshot of the real workout sheet for the dev server.

Download the Google Sheet as .xlsx (File > Download > Microsoft Excel) and run:
    .venv/bin/python devserver/xlsx2json.py log.xlsx devserver/data/snapshot.json
Keeps the tabs the backend reads (cycleN, rm calc, session log, hr, assistance plan) with cached
values, formulas and notes. The output holds your data, so devserver/data/ is gitignored.
"""

import datetime as dt
import json
import re
import sys

import openpyxl

KEEP = re.compile(r"^(cycle\d+|rm calc|session log|hr|assistance plan)$")


def encode(v):
    if isinstance(v, dt.datetime):
        return {"$date": v.replace(tzinfo=dt.timezone.utc).isoformat().replace("+00:00", "Z")}
    if isinstance(v, dt.date):
        return {"$date": dt.datetime(v.year, v.month, v.day, tzinfo=dt.timezone.utc).isoformat().replace("+00:00", "Z")}
    if isinstance(v, dt.time):
        return v.isoformat()
    return v


def main(src, dst):
    values_wb = openpyxl.load_workbook(src, data_only=True)
    formulas_wb = openpyxl.load_workbook(src, data_only=False)
    sheets = []
    for name in values_wb.sheetnames:
        if not KEEP.match(name):
            continue
        vs, fs = values_wb[name], formulas_wb[name]
        values, formulas, notes = [], {}, {}
        for row in vs.iter_rows():
            out = []
            for cell in row:
                out.append(encode(cell.value) if cell.value is not None else "")
                if cell.comment:
                    notes[f"{cell.row},{cell.column}"] = cell.comment.text
            while out and out[-1] == "":
                out.pop()
            values.append(out)
        for row in fs.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas[f"{cell.row},{cell.column}"] = cell.value
        sheets.append({"name": name, "values": values, "formulas": formulas, "notes": notes})
        print(f"{name}: {len(values)} rows, {len(formulas)} formulas", file=sys.stderr)
    with open(dst, "w") as f:
        json.dump({"tz": "Europe/Prague", "sheets": sheets, "props": {}}, f)


if __name__ == "__main__":
    main(*sys.argv[1:3])
